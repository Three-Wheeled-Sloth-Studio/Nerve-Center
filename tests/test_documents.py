from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from nerve_center.profile.documents import import_source_document
from nerve_center.profile.models import DocumentFormat


def test_imports_text_and_markdown_without_copying(tmp_path: Path) -> None:
    source = tmp_path / "resume.md"
    source.write_text("# Product Leader\n\nReduced cycle time by 80%.\n", encoding="utf-8")

    document = import_source_document(source)

    assert document.format is DocumentFormat.MARKDOWN
    assert document.source_path == str(source.resolve())
    assert document.segments[0].locator == "line:1"
    assert "80%" in document.analysis_text


def test_imports_docx(tmp_path: Path) -> None:
    source = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Data Product Manager")
    document.add_paragraph("Managed teams across five countries.")
    document.save(source)

    imported = import_source_document(source)

    assert imported.format is DocumentFormat.DOCX
    assert [item.locator for item in imported.segments] == ["paragraph:1", "paragraph:2"]


def test_imports_pdf(tmp_path: Path) -> None:
    source = tmp_path / "resume.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 720, "Built a five million dollar analytics product line.")
    pdf.save()

    imported = import_source_document(source)

    assert imported.format is DocumentFormat.PDF
    assert any("analytics product line" in item.text for item in imported.segments)
