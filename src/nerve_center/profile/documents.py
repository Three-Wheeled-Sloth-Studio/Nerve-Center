"""Read-only local resume document registration and text extraction."""

from __future__ import annotations

import hashlib
import re
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader

from nerve_center.profile.models import DocumentFormat, DocumentSegment, SourceDocument

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
_HORIZONTAL_SPACE = re.compile(r"[\t\f\v ]+")


class DocumentImportError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def import_source_document(path: Path) -> SourceDocument:
    resolved = path.expanduser().resolve(strict=True)
    if not resolved.is_file():
        raise DocumentImportError("NOT_A_FILE", "The selected resume source is not a file.")

    byte_size = resolved.stat().st_size
    if byte_size > MAX_DOCUMENT_BYTES:
        raise DocumentImportError(
            "DOCUMENT_TOO_LARGE",
            "Resume source files must be 25 MB or smaller.",
        )

    raw = resolved.read_bytes()
    document_format = _detect_format(resolved)
    segments = _extract_segments(raw, document_format)
    if not segments:
        raise DocumentImportError(
            "NO_EXTRACTABLE_TEXT",
            "No usable text could be extracted from the selected resume source.",
        )

    stat = resolved.stat()
    return SourceDocument(
        id=str(uuid4()),
        source_path=str(resolved),
        file_name=resolved.name,
        format=document_format,
        media_type=_media_type(document_format),
        sha256=hashlib.sha256(raw).hexdigest(),
        byte_size=byte_size,
        modified_at=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
        segments=segments,
    )


def _detect_format(path: Path) -> DocumentFormat:
    suffix = path.suffix.lower()
    mapping = {
        ".txt": DocumentFormat.TEXT,
        ".md": DocumentFormat.MARKDOWN,
        ".markdown": DocumentFormat.MARKDOWN,
        ".docx": DocumentFormat.DOCX,
        ".pdf": DocumentFormat.PDF,
    }
    try:
        return mapping[suffix]
    except KeyError as error:
        raise DocumentImportError(
            "UNSUPPORTED_DOCUMENT_FORMAT",
            "Supported resume formats are PDF, DOCX, Markdown, and plain text.",
        ) from error


def _extract_segments(raw: bytes, document_format: DocumentFormat) -> list[DocumentSegment]:
    if document_format in {DocumentFormat.TEXT, DocumentFormat.MARKDOWN}:
        return _line_segments(_decode_text(raw), "line")
    if document_format is DocumentFormat.DOCX:
        return _docx_segments(raw)
    if document_format is DocumentFormat.PDF:
        return _pdf_segments(raw)
    raise AssertionError(f"unsupported document format: {document_format}")


def _decode_text(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252")


def _line_segments(text: str, prefix: str) -> list[DocumentSegment]:
    result: list[DocumentSegment] = []
    for index, value in enumerate(text.splitlines(), start=1):
        cleaned = _clean_segment(value)
        if cleaned:
            result.append(DocumentSegment(locator=f"{prefix}:{index}", text=cleaned))
    return result


def _docx_segments(raw: bytes) -> list[DocumentSegment]:
    try:
        document = Document(BytesIO(raw))
    except Exception as error:
        raise DocumentImportError(
            "DOCX_PARSE_FAILED",
            "The DOCX file could not be read.",
        ) from error

    result: list[DocumentSegment] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        cleaned = _clean_segment(paragraph.text)
        if cleaned:
            result.append(DocumentSegment(locator=f"paragraph:{index}", text=cleaned))
    return result


def _pdf_segments(raw: bytes) -> list[DocumentSegment]:
    try:
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise DocumentImportError(
                "PDF_ENCRYPTED",
                "Encrypted PDF resume sources are not supported.",
            )
        result: list[DocumentSegment] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            result.extend(_line_segments(text, f"page:{page_index}:line"))
        return result
    except DocumentImportError:
        raise
    except Exception as error:
        raise DocumentImportError("PDF_PARSE_FAILED", "The PDF file could not be read.") from error


def _clean_segment(value: str) -> str:
    value = value.replace("\x00", "").strip()
    return _HORIZONTAL_SPACE.sub(" ", value)


def _media_type(document_format: DocumentFormat) -> str:
    return {
        DocumentFormat.TEXT: "text/plain",
        DocumentFormat.MARKDOWN: "text/markdown",
        DocumentFormat.DOCX: (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        DocumentFormat.PDF: "application/pdf",
    }[document_format]
